# -*- coding: utf-8 -*-
"""生成家庭牧场记录 Excel 模板（三个 Sheet，带示例行与格式）"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "家庭牧场记录模板.xlsx")

HEAD_FILL = PatternFill("solid", fgColor="4472C4")
HEAD_FONT = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="微软雅黑", size=11)
SAMPLE_FONT = Font(name="微软雅黑", size=11, italic=True, color="808080")
CENTER = Alignment(horizontal="center", vertical="center")

SHEETS = [
    ("卖牛羊记录", [
        ("日期", 12), ("品种", 10), ("阶段", 10), ("数量", 8),
        ("平均体重(kg)", 13), ("单价(元/斤)", 12), ("买家", 14), ("备注", 22),
    ], [
        ["2026-10-01", "黄牛", "育肥牛", 2, 480, 15.2, "贩子老王", "膘好加了一毛"],
        ["2026-10-03", "藏羊", "羊羔", 10, 35, 18.0, "合作社", "一批出栏"],
    ]),
    ("买饲料记录", [
        ("日期", 12), ("饲料种类", 12), ("数量(kg)", 10),
        ("单价(元/kg)", 12), ("供应商", 14), ("备注", 22),
    ], [
        ["2026-10-02", "玉米", 1000, 2.40, "粮站", "赊账月底结"],
        ["2026-10-05", "豆粕", 200, 3.35, "饲料店小李", "现金"],
    ]),
    ("贩子报价记录", [
        ("日期", 12), ("品种", 10), ("阶段", 10),
        ("报价(元/斤)", 12), ("报价人", 12), ("备注", 22),
    ], [
        ["2026-10-04", "黄牛", "育肥牛", 15.0, "贩子老王", "电话报价"],
        ["2026-10-06", "牦牛", "牛犊", 12.5, "贩子老马", "看牛后口头价"],
    ]),
]

wb = Workbook()
wb.remove(wb.active)
for name, cols, samples in SHEETS:
    ws = wb.create_sheet(name)
    for i, (title, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=title)
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, CENTER
        ws.column_dimensions[get_column_letter(i)].width = width
    for r, row in enumerate(samples, 2):
        for c_i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=c_i, value=v)
            c.font = SAMPLE_FONT
    ws.freeze_panes = "A2"
    # 说明行
    note = ws.cell(row=len(samples) + 3, column=1,
                   value="（灰色为示例行，照着格式往下记；日期写 2026-10-01 这种格式）")
    note.font = Font(name="微软雅黑", size=10, color="808080")

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("模板已生成:", OUT)
